"""Generic OHLC order-block strategy and Excel backtest report.

Default input: GOX25-Z25_5 min.parquet in the same folder as this script.

The script also accepts --input so it can be run from any machine/path.
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


@dataclass
class BacktestConfig:
    tick_size: float = 0.25
    tick_value: float = 25.0  # dollars earned/lost per tick per lot
    fees_per_round_trip: float = 2.0
    lots: int = 1
    initial_capital: float = 0.0
    entry_delay_bars: int = 1
    conservative_same_bar: bool = True
    minimum_move_ticks: float = 1.0
    maximum_move_ticks: float = 3.0
    partial_mode: str = "OFF"  # OFF, INTERVAL, or PERCENT
    partial_interval_ticks: float = 2.0
    trailing_sl_enabled: bool = False
    exclude_overnight_trades: bool = False


class Strategy:
    """Order-block backtest that infers its candle interval from the data."""

    def __init__(self, parameters, backtest_config, data, data_dict=None):
        self.parameters = parameters or {}
        self.backtest_config = backtest_config or {}
        self.data = data.copy()
        self.data_dict = data_dict or {}
        self.cfg = BacktestConfig(
            tick_size=float(self.parameters.get("Tick_Size", 0.25)),
            tick_value=float(self.parameters.get("Tick_Value", 25.0)),
            fees_per_round_trip=float(self.parameters.get("Fees", 2.0)),
            lots=int(self.parameters.get("Lots", 1)),
            initial_capital=float(self.parameters.get("Initial_Capital", 0.0)),
            entry_delay_bars=int(self.parameters.get("Entry_Delay_Bars", 1)),
            conservative_same_bar=bool(self.parameters.get("Conservative_Same_Bar", True)),
            minimum_move_ticks=float(self.parameters.get("Minimum_Move_Ticks", 1.0)),
            maximum_move_ticks=float(self.parameters.get("Maximum_Move_Ticks", 3.0)),
            partial_mode=str(self.parameters.get("Partial_Mode", "OFF")).upper(),
            partial_interval_ticks=float(self.parameters.get("Partial_Interval_Ticks", 2.0)),
            trailing_sl_enabled=bool(self.parameters.get("Trailing_SL_Enabled", False)),
            exclude_overnight_trades=bool(self.parameters.get("Exclude_Overnight_Trades", False)),
        )
        if self.cfg.minimum_move_ticks <= 0 or self.cfg.maximum_move_ticks < self.cfg.minimum_move_ticks:
            raise ValueError("Order-block tick range must satisfy 0 < minimum <= maximum")
        if self.cfg.lots < 1:
            raise ValueError("Lots must be at least 1")
        if self.cfg.partial_mode not in {"OFF", "INTERVAL", "PERCENT"}:
            raise ValueError("Partial mode must be OFF, INTERVAL, or PERCENT")
        if self.cfg.partial_mode == "INTERVAL" and self.cfg.partial_interval_ticks <= 0:
            raise ValueError("Partial-booking interval must be greater than zero")
        raw_percentages = self.parameters.get("Partial_Target_Percentages", [25, 50, 75])
        if isinstance(raw_percentages, str):
            raw_percentages = [value.strip() for value in raw_percentages.split(",") if value.strip()]
        self.partial_percentages = sorted({float(value) for value in raw_percentages})
        if self.cfg.partial_mode == "PERCENT" and not all(0 < value < 100 for value in self.partial_percentages):
            raise ValueError("Partial target percentages must be between 0 and 100")
        self.bar_delta = None
        self.overnight_trades_removed = 0

    def _clean(self) -> pd.DataFrame:
        df = self.data.copy()
        df.columns = [str(c).strip().lower() for c in df.columns]
        required = {"time", "open", "high", "low", "close"}
        missing = required.difference(df.columns)
        if missing:
            raise ValueError(f"Missing required parquet columns: {sorted(missing)}")
        df["time"] = pd.to_datetime(df["time"], errors="coerce")
        for col in ["open", "high", "low", "close"]:
            df[col] = pd.to_numeric(df[col], errors="coerce")
        df = (df.dropna(subset=["time", "open", "high", "low", "close"])
                .sort_values("time").drop_duplicates("time", keep="last").reset_index(drop=True))
        return df

    @staticmethod
    def _infer_bar_delta(df: pd.DataFrame) -> pd.Timedelta:
        """Infer the normal candle interval from the most frequent positive gap."""
        differences = df["time"].diff().dropna()
        differences = differences[differences > pd.Timedelta(0)]
        if differences.empty:
            raise ValueError("At least two valid candles are required to infer the interval")
        return differences.value_counts().idxmax()

    def _signals(self, df: pd.DataFrame) -> pd.DataFrame:
        tick = self.cfg.tick_size
        prev_open = df["open"].shift(1)
        prev_close = df["close"].shift(1)
        self.bar_delta = self._infer_bar_delta(df)
        contiguous = df["time"].diff().eq(self.bar_delta)
        same_open = np.isclose(df["open"], prev_close, atol=tick * 1e-6, rtol=0)
        previous_is_bearish = prev_open > prev_close
        previous_is_bullish = prev_open < prev_close
        bullish_move = df["close"] - prev_open
        bearish_move = prev_open - df["close"]
        df["bull_order_block"] = (
            contiguous
            & same_open
            & previous_is_bearish
            & (bullish_move >= self.cfg.minimum_move_ticks * tick)
            & (bullish_move <= self.cfg.maximum_move_ticks * tick)
        )
        df["bear_order_block"] = (
            contiguous
            & same_open
            & previous_is_bullish
            & (bearish_move >= self.cfg.minimum_move_ticks * tick)
            & (bearish_move <= self.cfg.maximum_move_ticks * tick)
        )
        return df

    def StrategyBuilder(self):
        df = self._signals(self._clean())
        trades = self._run_trades(df)
        if self.cfg.exclude_overnight_trades and not trades.empty:
            intraday = (
                pd.to_datetime(trades["entry_time"]).dt.normalize()
                == pd.to_datetime(trades["exit_time"]).dt.normalize()
            )
            self.overnight_trades_removed = int((~intraday).sum())
            trades = trades.loc[intraday].reset_index(drop=True)
            trades["cumulative_pnl"] = trades["net_pnl"].cumsum()
        equity = self._equity_curve(df, trades)
        stats = self._statistics(df, trades, equity)
        return df, trades, equity, stats

    def _run_trades(self, df: pd.DataFrame) -> pd.DataFrame:
        trades, position, pending = [], None, None
        trade_id = 0

        def close_position(pos, bar, index, exit_price, reason):
            """Close all remaining lots and combine them with any partial fill."""
            remaining = pos["remaining_lots"]
            final_ticks = (pos["direction"] * (exit_price - pos["entry_price"])
                           / self.cfg.tick_size * remaining)
            gross_ticks = pos["partial_gross_ticks"] + final_ticks
            gross_pnl = gross_ticks * self.cfg.tick_value
            fees = self.cfg.fees_per_round_trip * self.cfg.lots
            pos.update({
                "exit_time": bar.time, "exit_price": float(exit_price), "exit_reason": reason,
                "lots": self.cfg.lots, "remaining_lots_at_exit": remaining,
                "gross_ticks": gross_ticks, "gross_pnl": gross_pnl, "fees": fees,
                "net_pnl": gross_pnl - fees, "bars_held": index - pos["entry_index"] + 1,
                "stop_ticks": pos["risk_price"] / self.cfg.tick_size,
                "take_profit_ticks": pos["risk_price"] / self.cfg.tick_size,
                "final_stop_price": pos["stop_price"],
                "partial_bookings_count": len(pos["partial_fills"]),
                "partial_exit_details": " | ".join(
                    f"1 lot @ {fill['price']:.8g} ({fill['ticks']:.4g} ticks, {fill['time']})"
                    for fill in pos["partial_fills"]
                ) or "None",
            })
            pos.pop("partial_levels_ticks", None)
            pos.pop("next_partial_index", None)
            pos.pop("partial_fills", None)
            trades.append(pos)

        def partial_levels(risk_ticks):
            """Return ordered one-lot exit levels strictly below the final target."""
            if self.cfg.lots <= 1 or self.cfg.partial_mode == "OFF":
                return []
            if self.cfg.partial_mode == "PERCENT":
                levels = [risk_ticks * percent / 100 for percent in self.partial_percentages]
            else:
                levels = [self.cfg.partial_interval_ticks * number
                          for number in range(1, self.cfg.lots)]
            return [level for level in levels if level < risk_ticks - 1e-9][:self.cfg.lots - 1]

        def book_reached_partials(pos, bar, favorable_ticks):
            """Book one lot at every newly crossed level, in price order."""
            while (pos["remaining_lots"] > 1
                   and pos["next_partial_index"] < len(pos["partial_levels_ticks"])):
                level = pos["partial_levels_ticks"][pos["next_partial_index"]]
                if favorable_ticks + 1e-9 < level:
                    break
                price = pos["entry_price"] + pos["direction"] * level * self.cfg.tick_size
                pos["partial_fills"].append({"time": bar.time, "price": price, "ticks": level})
                pos["partial_lots_booked"] += 1
                pos["remaining_lots"] -= 1
                pos["partial_exit_time"] = bar.time
                pos["partial_exit_price"] = price
                pos["partial_profit_ticks"] = level
                pos["partial_gross_ticks"] += level
                pos["partial_gross_pnl"] = pos["partial_gross_ticks"] * self.cfg.tick_value
                pos["next_partial_index"] += 1

                # Distribute all booked profit ticks over the lots still at risk.
                adjustment_ticks = pos["partial_gross_ticks"] / pos["remaining_lots"]
                pos["partial_stop_adjustment_ticks"] = adjustment_ticks
                adjusted_stop = (pos["initial_stop_price"] + pos["direction"]
                                 * adjustment_ticks * self.cfg.tick_size)
                pos["stop_price"] = (max(pos["stop_price"], adjusted_stop)
                                     if pos["direction"] == 1
                                     else min(pos["stop_price"], adjusted_stop))

        for i, bar in df.iterrows():
            # Signals execute at the next valid, consecutive candle open.
            if pending and i == pending["entry_index"]:
                entry = float(bar.open)
                direction = pending["direction"]
                stop = pending["stop"]
                risk = entry - stop if direction == 1 else stop - entry
                if risk > 0:
                    risk_ticks = risk / self.cfg.tick_size
                    position = {
                        "trade_id": trade_id + 1, "side": "LONG" if direction == 1 else "SHORT",
                        "direction": direction, "signal_time": pending["signal_time"],
                        "entry_time": bar.time, "entry_price": entry,
                        "initial_stop_price": stop, "stop_price": stop,
                        # The final target is always fixed at 1:1 reward-to-risk.
                        "target_price": entry + direction * risk,
                        "risk_price": risk,
                        "entry_index": i,
                        "remaining_lots": self.cfg.lots,
                        "partial_lots_booked": 0,
                        "partial_exit_time": pd.NaT,
                        "partial_exit_price": np.nan,
                        "partial_profit_ticks": np.nan,
                        "partial_gross_ticks": 0.0,
                        "partial_gross_pnl": 0.0,
                        "partial_stop_adjustment_ticks": 0.0,
                        "partial_levels_ticks": partial_levels(risk_ticks),
                        "next_partial_index": 0,
                        "partial_fills": [],
                        "trailing_stop_updates": 0,
                        "max_favorable_ticks": 0.0,
                    }
                    trade_id += 1
                pending = None

            if position:
                d, stop, target = position["direction"], position["stop_price"], position["target_price"]
                reason = exit_price = None
                # Opening gaps fill at the bar open, not at an unavailable stop/target quote.
                if (d == 1 and bar.open <= stop) or (d == -1 and bar.open >= stop):
                    exit_price, reason = float(bar.open), "SL_GAP"
                elif (d == 1 and bar.open >= target) or (d == -1 and bar.open <= target):
                    exit_price, reason = float(bar.open), "TP_GAP"
                else:
                    stop_hit = bar.low <= stop if d == 1 else bar.high >= stop
                    target_hit = bar.high >= target if d == 1 else bar.low <= target
                    if stop_hit and target_hit:
                        exit_price, reason = ((stop, "SL_SAME_BAR") if self.cfg.conservative_same_bar
                                              else (target, "TP_SAME_BAR"))
                    elif stop_hit:
                        exit_price, reason = stop, "SL"
                    elif target_hit:
                        book_reached_partials(position, bar, position["risk_price"] / self.cfg.tick_size)
                        exit_price, reason = target, "TP"
                if exit_price is not None:
                    close_position(position, bar, i, float(exit_price), reason)
                    position = None

                if position:
                    # Trailing changes apply from the following candle: every 2 favorable
                    # ticks advances the stop 1 tick, without moving the fixed final target.
                    favorable_ticks = ((bar.high - position["entry_price"]) / self.cfg.tick_size
                                       if d == 1 else
                                       (position["entry_price"] - bar.low) / self.cfg.tick_size)
                    position["max_favorable_ticks"] = max(position["max_favorable_ticks"], favorable_ticks)
                    book_reached_partials(position, bar, position["max_favorable_ticks"])
                    if self.cfg.trailing_sl_enabled:
                        updates = max(0, int(np.floor((position["max_favorable_ticks"] + 1e-9) / 2.0)))
                        trailed_stop = position["initial_stop_price"] + d * updates * self.cfg.tick_size
                        position["stop_price"] = (max(position["stop_price"], trailed_stop)
                                                  if d == 1 else min(position["stop_price"], trailed_stop))
                        position["trailing_stop_updates"] = updates

            # One position at a time; a new signal is queued only while flat.
            if position is None and pending is None and i + 1 < len(df):
                next_is_contiguous = (
                    df.at[i + 1, "time"] - bar.time
                    == self.bar_delta
                )
                if next_is_contiguous and (bar.bull_order_block or bar.bear_order_block):
                    direction = 1 if bar.bull_order_block else -1
                    # Long: one tick below signal's previous candle low. Short: one tick above its high.
                    reference = df.iloc[i - 1]
                    stop = float(reference.low - self.cfg.tick_size if direction == 1
                                 else reference.high + self.cfg.tick_size)
                    pending = {"direction": direction, "stop": stop,
                               "signal_time": bar.time, "entry_index": i + self.cfg.entry_delay_bars}

        if position:
            bar = df.iloc[-1]
            close_position(position, bar, len(df) - 1, float(bar.close), "END_OF_DATA")
        result = pd.DataFrame(trades)
        if not result.empty:
            result["cumulative_pnl"] = result["net_pnl"].cumsum()
        return result

    def _equity_curve(self, df, trades):
        curve = df[["time"]].copy()
        realized = trades.groupby("exit_time")["net_pnl"].sum() if not trades.empty else pd.Series(dtype=float)
        curve["realized_pnl"] = curve["time"].map(realized).fillna(0.0)
        curve["equity"] = self.cfg.initial_capital + curve["realized_pnl"].cumsum()
        curve["running_peak"] = curve["equity"].cummax()
        curve["drawdown"] = curve["equity"] - curve["running_peak"]
        return curve

    @staticmethod
    def _max_streak(values, winning):
        best = current = 0
        for value in values:
            match = value > 0 if winning else value < 0
            current = current + 1 if match else 0
            best = max(best, current)
        return best

    def _statistics(self, df, trades, equity):
        pnl = trades["net_pnl"] if not trades.empty else pd.Series(dtype=float)
        returns = equity.set_index("time")["realized_pnl"].resample("D").sum()
        sharpe = np.sqrt(252) * returns.mean() / returns.std(ddof=1) if returns.std(ddof=1) > 0 else np.nan
        wins, losses = int((pnl > 0).sum()), int((pnl < 0).sum())
        winning_longs = int(((trades["side"] == "LONG") & (pnl > 0)).sum()) if not trades.empty else 0
        winning_shorts = int(((trades["side"] == "SHORT") & (pnl > 0)).sum()) if not trades.empty else 0
        sl_hits = int(trades["exit_reason"].str.startswith("SL").sum()) if not trades.empty else 0
        return pd.DataFrame([{
            "data_start": df.time.min(), "data_end": df.time.max(), "valid_bars": len(df),
            "bull_order_blocks": int(df.bull_order_block.sum()), "bear_order_blocks": int(df.bear_order_block.sum()),
            "total_trades": len(trades), "profitable_trades": wins, "losing_trades": losses,
            "winning_long_trades": winning_longs, "winning_short_trades": winning_shorts,
            "breakeven_trades": int((pnl == 0).sum()), "win_rate_pct": 100*wins/len(trades) if len(trades) else np.nan,
            "stop_loss_hits": sl_hits, "take_profit_hits": int(trades.exit_reason.str.startswith("TP").sum()) if len(trades) else 0,
            "gross_pnl": trades.gross_pnl.sum() if len(trades) else 0.0,
            "total_fees": trades.fees.sum() if len(trades) else 0.0, "net_pnl": pnl.sum(),
            "max_trade_profit": pnl.max() if len(trades) else np.nan,
            "max_trade_loss": pnl.min() if len(trades) else np.nan,
            "max_profit_streak": self._max_streak(pnl, True), "max_loss_streak": self._max_streak(pnl, False),
            "max_drawdown": equity.drawdown.min(), "daily_sharpe_ratio": sharpe,
            "tick_size": self.cfg.tick_size, "tick_value_dollars": self.cfg.tick_value,
            "fees_per_round_trip": self.cfg.fees_per_round_trip, "lots": self.cfg.lots,
            "minimum_move_ticks": self.cfg.minimum_move_ticks,
            "maximum_move_ticks": self.cfg.maximum_move_ticks,
            "fixed_tp_sl_ratio": "1:1",
            "partial_mode": self.cfg.partial_mode,
            "partial_interval_ticks": self.cfg.partial_interval_ticks if self.cfg.partial_mode == "INTERVAL" else np.nan,
            "partial_target_percentages": ", ".join(f"{value:g}%" for value in self.partial_percentages) if self.cfg.partial_mode == "PERCENT" else "",
            "partial_bookings": int(trades["partial_bookings_count"].sum()) if len(trades) else 0,
            "trailing_sl_enabled": self.cfg.trailing_sl_enabled,
            "exclude_overnight_trades": self.cfg.exclude_overnight_trades,
            "overnight_trades_removed": self.overnight_trades_removed,
            "inferred_bar_minutes": self.bar_delta.total_seconds() / 60,
        }])


def save_excel_report(output_file, stats, trades, equity, order_blocks, market_chart_data):
    """Save all backtest results in one formatted Excel workbook."""
    summary = stats.T.reset_index()
    summary.columns = ["Metric", "Value"]

    sheets = {
        "Summary": summary,
        "Trades": trades,
        "Equity Curve": equity,
        "Market Chart": market_chart_data,
        "Order Blocks": order_blocks,
    }

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        for sheet_name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=sheet_name, index=False)

        workbook = writer.book
        dark_blue = "1F4E78"
        light_blue = "D9EAF7"
        green = "C6EFCE"
        red = "FFC7CE"

        for sheet_name, frame in sheets.items():
            ws = workbook[sheet_name]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            ws.sheet_view.showGridLines = False

            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor=dark_blue)
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 24

            if len(frame) > 0 and len(frame.columns) > 0:
                table_name = "Table" + sheet_name.replace(" ", "")
                table = Table(displayName=table_name, ref=ws.dimensions)
                table.tableStyleInfo = TableStyleInfo(
                    name="TableStyleMedium2",
                    showFirstColumn=False,
                    showLastColumn=False,
                    showRowStripes=True,
                    showColumnStripes=False,
                )
                ws.add_table(table)

            for column_number, column_name in enumerate(frame.columns, start=1):
                values = [str(column_name)] + [
                    "" if pd.isna(value) else str(value)
                    for value in frame.iloc[:, column_number - 1].head(1000)
                ]
                width = min(max(len(value) for value in values) + 2, 24)
                ws.column_dimensions[get_column_letter(column_number)].width = max(width, 12)

                name = str(column_name).lower()
                if "time" in name or "date" in name:
                    for cell in ws[get_column_letter(column_number)][1:]:
                        cell.number_format = "yyyy-mm-dd hh:mm"
                elif "price" in name or name in {"open", "high", "low", "close", "risk_price"}:
                    for cell in ws[get_column_letter(column_number)][1:]:
                        cell.number_format = "0.00"
                elif "pnl" in name or "fees" in name or "equity" in name or "drawdown" in name:
                    for cell in ws[get_column_letter(column_number)][1:]:
                        cell.number_format = '$#,##0.00;[Red]-$#,##0.00'

        summary_ws = workbook["Summary"]
        summary_ws.column_dimensions["A"].width = 28
        summary_ws.column_dimensions["B"].width = 24
        for row in range(2, summary_ws.max_row + 1):
            summary_ws.cell(row, 1).font = Font(bold=True, color=dark_blue)
            summary_ws.cell(row, 1).fill = PatternFill("solid", fgColor=light_blue)

        trades_ws = workbook["Trades"]
        if trades_ws.max_row > 1 and "net_pnl" in trades.columns:
            pnl_col = trades.columns.get_loc("net_pnl") + 1
            pnl_range = f"{get_column_letter(pnl_col)}2:{get_column_letter(pnl_col)}{trades_ws.max_row}"
            trades_ws.conditional_formatting.add(
                pnl_range, CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=green))
            )
            trades_ws.conditional_formatting.add(
                pnl_range, CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor=red))
            )

        # Trade-by-trade cumulative net P&L chart.
        if trades_ws.max_row > 1 and "cumulative_pnl" in trades.columns:
            trade_id_col = trades.columns.get_loc("trade_id") + 1
            cumulative_col = trades.columns.get_loc("cumulative_pnl") + 1
            pnl_chart = LineChart()
            pnl_chart.title = "Cumulative Net P&L"
            pnl_chart.y_axis.title = "Cumulative P&L ($)"
            pnl_chart.x_axis.title = "Trade Number"
            pnl_chart.style = 13
            pnl_chart.height = 9
            pnl_chart.width = 18
            pnl_chart.add_data(
                Reference(trades_ws, min_col=cumulative_col, min_row=1, max_row=trades_ws.max_row),
                titles_from_data=True,
            )
            pnl_chart.set_categories(
                Reference(trades_ws, min_col=trade_id_col, min_row=2, max_row=trades_ws.max_row)
            )
            pnl_chart.series[0].graphicalProperties.line.solidFill = "1F4E78"
            pnl_chart.series[0].graphicalProperties.line.width = 24000
            trades_ws.add_chart(pnl_chart, f"{get_column_letter(len(trades.columns) + 2)}2")

        equity_ws = workbook["Equity Curve"]
        if equity_ws.max_row > 1:
            time_col = equity.columns.get_loc("time") + 1
            equity_col = equity.columns.get_loc("equity") + 1
            chart = LineChart()
            chart.title = "GOX25-Z25 Equity Curve"
            chart.y_axis.title = "Equity ($)"
            chart.x_axis.title = "Time"
            chart.style = 13
            chart.height = 9
            chart.width = 18
            chart.add_data(
                Reference(equity_ws, min_col=equity_col, min_row=1, max_row=equity_ws.max_row),
                titles_from_data=True,
            )
            chart.set_categories(
                Reference(equity_ws, min_col=time_col, min_row=2, max_row=equity_ws.max_row)
            )
            equity_ws.add_chart(chart, "G2")

        # Entire GOX25-Z25 close history with order blocks and executions.
        market_ws = workbook["Market Chart"]
        if market_ws.max_row > 1:
            market_chart = LineChart()
            market_chart.title = "GOX25-Z25 Full History: Order Blocks and Trades"
            market_chart.y_axis.title = "Spread Price"
            market_chart.x_axis.title = "Time"
            market_chart.style = 13
            market_chart.height = 15
            market_chart.width = 30
            time_col = market_chart_data.columns.get_loc("time") + 1
            close_col = market_chart_data.columns.get_loc("close") + 1
            market_chart.add_data(
                Reference(market_ws, min_col=close_col, min_row=1, max_row=market_ws.max_row),
                titles_from_data=True,
            )
            market_chart.set_categories(
                Reference(market_ws, min_col=time_col, min_row=2, max_row=market_ws.max_row)
            )
            market_chart.series[0].graphicalProperties.line.solidFill = "666666"
            market_chart.series[0].graphicalProperties.line.width = 12000

            marker_settings = {
                "bull_order_block_price": ("triangle", "00A65A"),
                "bear_order_block_price": ("triangle", "D62728"),
                "long_entry_price": ("diamond", "0070C0"),
                "short_entry_price": ("diamond", "C00000"),
                "exit_price": ("circle", "7030A0"),
            }
            for column_name, (symbol, color) in marker_settings.items():
                column_number = market_chart_data.columns.get_loc(column_name) + 1
                market_chart.add_data(
                    Reference(market_ws, min_col=column_number, min_row=1, max_row=market_ws.max_row),
                    titles_from_data=True,
                )
                series = market_chart.series[-1]
                series.graphicalProperties.line.noFill = True
                series.marker.symbol = symbol
                series.marker.size = 6
                series.marker.graphicalProperties.solidFill = color
                series.marker.graphicalProperties.line.solidFill = color

            market_chart.legend.position = "b"
            market_ws.add_chart(market_chart, f"{get_column_letter(len(market_chart_data.columns) + 2)}2")


def main():
    script_folder = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--input",
        default=str(script_folder / "GOX25-Z25_5 min.parquet"),
        help="Parquet file path. By default, the file is read from this script's folder.",
    )
    parser.add_argument(
        "--output-dir",
        default=str(script_folder / "GOX25_Z25_backtest_results"),
    )
    parser.add_argument("--tick-value", type=float, default=25.0,
                        help="Dollar value of one 0.25 tick per lot (default: $25).")
    args = parser.parse_args()
    data = pd.read_parquet(args.input)
    strategy = Strategy({"Tick_Size": 0.25, "Tick_Value": args.tick_value, "Fees": 2, "Lots": 1}, {}, data)
    bars, trades, equity, stats = strategy.StrategyBuilder()
    out = Path(args.output_dir)
    out.mkdir(parents=True, exist_ok=True)
    order_blocks = bars.loc[
        bars.bull_order_block | bars.bear_order_block,
        ["time", "open", "high", "low", "close", "bull_order_block", "bear_order_block"],
    ].copy()
    market_chart_data = bars[["time", "open", "high", "low", "close"]].copy()
    market_chart_data["bull_order_block_price"] = np.where(
        bars["bull_order_block"], bars["low"], np.nan
    )
    market_chart_data["bear_order_block_price"] = np.where(
        bars["bear_order_block"], bars["high"], np.nan
    )
    market_chart_data["long_entry_price"] = np.nan
    market_chart_data["short_entry_price"] = np.nan
    market_chart_data["exit_price"] = np.nan
    if not trades.empty:
        long_entries = trades.loc[trades["side"] == "LONG"].set_index("entry_time")["entry_price"]
        short_entries = trades.loc[trades["side"] == "SHORT"].set_index("entry_time")["entry_price"]
        exits = trades.set_index("exit_time")["exit_price"]
        market_chart_data["long_entry_price"] = market_chart_data["time"].map(long_entries)
        market_chart_data["short_entry_price"] = market_chart_data["time"].map(short_entries)
        market_chart_data["exit_price"] = market_chart_data["time"].map(exits)
    output_file = out / "GOX25_Z25_Backtest_Report.xlsx"
    save_excel_report(output_file, stats, trades, equity, order_blocks, market_chart_data)
    print(stats.to_string(index=False))
    print(f"\nExcel report saved to: {output_file.resolve()}")


if __name__ == "__main__":
    main()
